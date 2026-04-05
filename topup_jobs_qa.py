import json
import logging
import os
import random
import re
import time
from copy import deepcopy
from datetime import datetime, timezone
from urllib.parse import urlparse, urlunparse

import openpyxl
import requests
from bs4 import BeautifulSoup

try:
    import yaml
except Exception:  # pragma: no cover
    yaml = None

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
RESULTS_DIR_DEFAULT = os.path.join(BASE_DIR, 'results')
CONFIG_PATH_DEFAULT = os.path.join(BASE_DIR, 'config.yaml')

DEFAULT_CONFIG = {
    'target': 150,
    'user_agent': 'Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36',
    'request': {
        'timeout_seconds': 20,
        'retries': 3,
        'backoff_seconds': 1.0,
        'jitter_seconds': 0.3,
    },
    'filters': {
        'role_keywords': ['qa', 'quality assurance', 'sdet', 'test', 'testing', 'automation'],
        'qa_keywords': ['qa', 'quality assurance', 'sdet', 'test', 'testing', 'automation'],
        'stack_keywords': ['python', 'playwright', 'cypress', 'api', 'ci/cd', 'pytest', 'selenium'],
        'remote_keywords': ['remote'],
        'us_keywords': ['united states', 'usa', 'us', 'u.s.', 'us only'],
        'unavailable_phrases': [
            'job is no longer available',
            'position has been filled',
            'no longer accepting applications',
            'this job has been removed',
            'job posting has expired',
            '404 not found',
            'page not found',
            'this position has been closed',
            'this posting is no longer available',
        ],
        'blocked_phrases': [
            'security check',
            'captcha',
            'verify you are a human',
            'access denied',
            'cloudflare ray id',
        ],
        'closed_url_patterns': [
            '/404',
            '/not-found',
            '/job-not-found',
            '/positions/closed',
        ],
    },
    'scoring': {
        'base_qa': 2,
        'stack_each': 1,
        'us_bonus': 1,
    },
    'sources': {
        'builtin_max_pages': 20,
        'themuse_max_pages': 120,
        'arbeitnow_max_pages': 20,
    },
    'boards': {
        'greenhouse': ['greenhouse', 'stripe', 'reddit', 'hubspot', 'asana', 'coursera', 'intercom'],
        'lever': ['lever', 'netflix', 'airtable', 'motive', 'scaleai', 'atlassian'],
    },
}

COLUMNS = [
    'Date Found',
    'Job Title',
    'Company',
    'Location',
    'Job URL',
    'Date Posted',
    'Min Salary',
    'Max Salary',
    'Employment Type',
    'Notes',
    'Application Status',
    'Applied Date',
    'Last Checked',
    'Link Status',
    'Job Key',
]


def deep_merge(base, override):
    out = deepcopy(base)
    for k, v in (override or {}).items():
        if isinstance(v, dict) and isinstance(out.get(k), dict):
            out[k] = deep_merge(out[k], v)
        else:
            out[k] = v
    return out


def load_config(path):
    cfg = deepcopy(DEFAULT_CONFIG)
    if os.path.exists(path) and yaml is not None:
        with open(path, 'r', encoding='utf-8') as f:
            user_cfg = yaml.safe_load(f) or {}
        cfg = deep_merge(cfg, user_cfg)
    return cfg


def load_yaml(path):
    if not os.path.exists(path) or yaml is None:
        return {}
    with open(path, 'r', encoding='utf-8') as f:
        return yaml.safe_load(f) or {}


def ensure_profile(profile_path, resume_dir):
    existing = load_yaml(profile_path)
    if existing:
        return existing
    try:
        import profile_builder
        profile = profile_builder.build_profile(resume_dir, profile_path)
        return profile or {}
    except Exception:
        return {}


CONFIG_PATH = os.getenv('CONFIG_FILE', CONFIG_PATH_DEFAULT)
KEYWORDS_PATH = os.getenv('KEYWORDS_FILE', os.path.join(BASE_DIR, 'keywords.yaml'))
CONFIG = load_config(CONFIG_PATH)

kw_cfg = load_yaml(KEYWORDS_PATH)
if isinstance(kw_cfg, dict):
    role_terms = kw_cfg.get('role_keywords')
    stack_terms = kw_cfg.get('stack_keywords')
    if isinstance(role_terms, list) and role_terms:
        CONFIG.setdefault('filters', {})['role_keywords'] = role_terms
        CONFIG['filters']['qa_keywords'] = role_terms
    if isinstance(stack_terms, list) and stack_terms:
        CONFIG.setdefault('filters', {})['stack_keywords'] = stack_terms

TARGET = int(os.getenv('TARGET', str(CONFIG['target'])))
RESULTS_DIR = os.getenv('RESULTS_DIR', RESULTS_DIR_DEFAULT)
RESUME_DIR = os.getenv('RESUME_DIR', os.path.join(BASE_DIR, 'resume'))
PROFILE_PATH = os.getenv('PROFILE_FILE', os.path.join(RESULTS_DIR, 'profile.yaml'))
FILE = os.getenv('JOBS_XLSX', os.path.join(RESULTS_DIR, 'jobs_qa.xlsx'))
JSON_FILE = os.getenv('JOBS_JSON', os.path.join(RESULTS_DIR, 'jobs_qa.json'))
PRUNE_CLOSED = os.getenv('PRUNE_CLOSED', '1').strip().lower() not in {'0', 'false', 'no'}
HEADERS = {'User-Agent': CONFIG['user_agent']}



def build_terms_re(terms):
    escaped = [re.escape(t) for t in terms if str(t).strip()]
    if not escaped:
        return re.compile(r'$^')
    return re.compile('(' + '|'.join(escaped) + ')', re.I)


def build_single_word_re(terms):
    escaped = [re.escape(t) for t in terms if str(t).strip()]
    if not escaped:
        return re.compile(r'$^')
    return re.compile(r'\b(' + '|'.join(escaped) + r')\b', re.I)


def _filters_from_config(config):
    filters = config.get('filters', {})
    role_terms = filters.get('role_keywords') or filters.get('qa_keywords') or []
    return {
        'role_keywords': role_terms,
        'stack_keywords': filters.get('stack_keywords', []),
        'remote_keywords': filters.get('remote_keywords', ['remote']),
        'us_keywords': filters.get('us_keywords', []),
        'unavailable_phrases': filters.get('unavailable_phrases', []),
        'blocked_phrases': filters.get('blocked_phrases', []),
        'closed_url_patterns': filters.get('closed_url_patterns', []),
    }


def _merge_profile_filters(base_filters, profile):
    out = deepcopy(base_filters)
    if not profile:
        return out
    for key in ['role_keywords', 'stack_keywords', 'remote_keywords', 'us_keywords']:
        vals = profile.get(key)
        if isinstance(vals, list) and vals:
            out[key] = list(dict.fromkeys([str(v).strip() for v in vals if str(v).strip()]))
    return out


def _slugify_builtin_term(term):
    t = re.sub(r'[^a-z0-9]+', '-', str(term).lower()).strip('-')
    return re.sub(r'-+', '-', t)


def _build_builtin_terms(profile=None):
    if profile:
        terms = (profile.get('title_keywords', []) + profile.get('role_keywords', []))
        out = []
        seen = set()
        for t in terms:
            slug = _slugify_builtin_term(t)
            if len(slug) < 2 or slug in seen:
                continue
            seen.add(slug)
            out.append(slug)
        if out:
            return out[:16]
    return [
        'qa', 'qa-automation-engineer', 'quality-assurance', 'quality-assurance-engineer',
        'qa-engineer', 'sdet', 'test-engineer', 'software-test-engineer', 'test-automation-engineer',
        'automation-test-engineer', 'quality-engineer'
    ]


def apply_runtime_filters(filters):
    global ROLE_RE, STACK_RE, REMOTE_RE, US_RE, UNAVAILABLE, BLOCKED, CLOSED_URL_PATTERNS
    ROLE_RE = build_terms_re(filters['role_keywords'])
    STACK_RE = build_terms_re(filters['stack_keywords'])
    REMOTE_RE = build_single_word_re(filters['remote_keywords'])
    US_RE = build_terms_re(filters['us_keywords'])
    UNAVAILABLE = [x.lower() for x in filters['unavailable_phrases']]
    BLOCKED = [x.lower() for x in filters['blocked_phrases']]
    CLOSED_URL_PATTERNS = [x.lower() for x in filters['closed_url_patterns']]


apply_runtime_filters(_filters_from_config(CONFIG))
GREENHOUSE_BOARDS = CONFIG['boards']['greenhouse']
LEVER_BOARDS = CONFIG['boards']['lever']
BUILTIN_TERMS = _build_builtin_terms()


RUN_STATS = {
    'run_started_utc': datetime.now(timezone.utc).isoformat(),
    'profile': {},
    'source_counts': {},
    'validation': {'ok': 0, 'bad': 0, 'reasons': {}},
    'dedupe': {},
}



def setup_logging():
    log_dir = os.path.join(BASE_DIR, 'logs')
    os.makedirs(log_dir, exist_ok=True)
    run_id = datetime.now().strftime('%Y%m%d_%H%M%S')
    log_path = os.path.join(log_dir, f'run_{run_id}.log')

    logger = logging.getLogger('jobs')
    logger.setLevel(logging.INFO)
    logger.handlers = []

    fmt = logging.Formatter('%(asctime)s level=%(levelname)s %(message)s')

    sh = logging.StreamHandler()
    sh.setFormatter(fmt)
    logger.addHandler(sh)

    fh = logging.FileHandler(log_path, encoding='utf-8')
    fh.setFormatter(fmt)
    logger.addHandler(fh)

    return logger, run_id, log_path


LOGGER, RUN_ID, LOG_PATH = setup_logging()



def log_event(event, **kwargs):
    msg = f'event={event}'
    if kwargs:
        pairs = [f'{k}={json.dumps(v, ensure_ascii=False)}' for k, v in kwargs.items()]
        msg = msg + ' ' + ' '.join(pairs)
    LOGGER.info(msg)



def norm(url: str) -> str:
    p = urlparse((url or '').strip())
    return urlunparse(('https', p.netloc.lower().replace('www.', ''), re.sub(r'/+$', '', p.path), '', '', '')).lower()



def today_str():
    return time.strftime('%Y-%m-%d')


def job_key_from_values(url: str, title: str, company: str, location: str):
    u = norm(url)
    if u:
        return u
    t = re.sub(r'\s+', ' ', str(title or '').strip().lower())
    c = re.sub(r'\s+', ' ', str(company or '').strip().lower())
    l = re.sub(r'\s+', ' ', str(location or '').strip().lower())
    return f'{t}|{c}|{l}'


def ensure_schema(ws):
    header = [c.value for c in ws[1]]
    h = {str(v).strip(): i for i, v in enumerate(header) if v}
    missing = [c for c in COLUMNS if c not in h]
    if not missing:
        return False
    for col in missing:
        ws.cell(row=1, column=ws.max_column + 1, value=col)
    return True


def hydrate_existing_rows(ws, h):
    changed = False
    for row_idx in range(2, ws.max_row + 1):
        title = str(ws.cell(row=row_idx, column=h['Job Title'] + 1).value or '').strip()
        company = str(ws.cell(row=row_idx, column=h['Company'] + 1).value or '').strip()
        location = str(ws.cell(row=row_idx, column=h['Location'] + 1).value or '').strip()
        url = str(ws.cell(row=row_idx, column=h['Job URL'] + 1).value or '').strip()
        if not any([title, company, location, url]):
            continue

        app_cell = ws.cell(row=row_idx, column=h['Application Status'] + 1)
        if not str(app_cell.value or '').strip():
            app_cell.value = 'new'
            changed = True

        link_cell = ws.cell(row=row_idx, column=h['Link Status'] + 1)
        if not str(link_cell.value or '').strip():
            link_cell.value = 'active'
            changed = True

        last_cell = ws.cell(row=row_idx, column=h['Last Checked'] + 1)
        if not str(last_cell.value or '').strip():
            last_cell.value = today_str()
            changed = True

        key_cell = ws.cell(row=row_idx, column=h['Job Key'] + 1)
        if not str(key_cell.value or '').strip():
            key_cell.value = job_key_from_values(url, title, company, location)
            changed = True
    return changed


def is_unavailable_html(text: str) -> bool:
    t = (text or '').lower()
    return any(x in t for x in UNAVAILABLE)



def is_blocked_html(text: str) -> bool:
    t = (text or '').lower()
    return any(x in t for x in BLOCKED)



def parse_salary_range(salary_text: str):
    text = str(salary_text or '').strip()
    if not text:
        return 'Not specified', 'Not specified'
    m = re.search(r'\$\s*([\d,]+(?:\.\d+)?)\s*(?:-|–|to)\s*\$\s*([\d,]+(?:\.\d+)?)', text, re.I)
    if m:
        return f"${m.group(1)}", f"${m.group(2)}"
    m1 = re.search(r'\$\s*([\d,]+(?:\.\d+)?)', text)
    if m1:
        return f"${m1.group(1)}", f"${m1.group(1)}"
    return 'Not specified', 'Not specified'



def request_with_retries(session, url, timeout=None):
    req_cfg = CONFIG['request']
    timeout = timeout or req_cfg['timeout_seconds']
    retries = int(req_cfg['retries'])
    backoff = float(req_cfg['backoff_seconds'])
    jitter = float(req_cfg['jitter_seconds'])

    last_err = None
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(url, timeout=timeout, allow_redirects=True)
            if resp.status_code in {429, 500, 502, 503, 504} and attempt < retries:
                sleep_s = backoff * attempt + random.uniform(0, jitter)
                time.sleep(sleep_s)
                continue
            return resp
        except Exception as e:  # pragma: no cover
            last_err = e
            if attempt < retries:
                sleep_s = backoff * attempt + random.uniform(0, jitter)
                time.sleep(sleep_s)
    if last_err:
        log_event('request_failed', url=url, error=str(last_err))
    return None



def validate_url(session, url: str):
    resp = request_with_retries(session, url)
    if resp is None:
        RUN_STATS['validation']['bad'] += 1
        RUN_STATS['validation']['reasons']['network_error'] = RUN_STATS['validation']['reasons'].get('network_error', 0) + 1
        return False, None, None, None, 'network_error'

    final_url = resp.url
    status = resp.status_code
    html = resp.text or ''
    final_url_norm = norm(final_url)

    if status >= 400:
        RUN_STATS['validation']['bad'] += 1
        reason = f'http_{status}'
        RUN_STATS['validation']['reasons'][reason] = RUN_STATS['validation']['reasons'].get(reason, 0) + 1
        return False, status, final_url, '', reason

    if any(x in final_url_norm for x in CLOSED_URL_PATTERNS):
        RUN_STATS['validation']['bad'] += 1
        RUN_STATS['validation']['reasons']['closed_url_pattern'] = RUN_STATS['validation']['reasons'].get('closed_url_pattern', 0) + 1
        return False, status, final_url, html, 'closed_url_pattern'

    if is_blocked_html(html):
        RUN_STATS['validation']['bad'] += 1
        RUN_STATS['validation']['reasons']['blocked_page'] = RUN_STATS['validation']['reasons'].get('blocked_page', 0) + 1
        return False, status, final_url, html, 'blocked_page'

    if is_unavailable_html(html):
        RUN_STATS['validation']['bad'] += 1
        RUN_STATS['validation']['reasons']['closed_phrase'] = RUN_STATS['validation']['reasons'].get('closed_phrase', 0) + 1
        return False, status, final_url, html, 'closed_phrase'

    RUN_STATS['validation']['ok'] += 1
    return True, status, final_url, html, 'ok'



def looks_like_us_remote(text: str, location: str = '') -> bool:
    hay = f'{text} {location}'.lower()
    remote_ok = bool(REMOTE_RE.search(hay))
    us_ok = bool(US_RE.search(hay))
    return remote_ok and us_ok



def normalize_location(location: str) -> str:
    l = (location or '').strip()
    if not l:
        return 'Remote'
    return l



def score_role(text: str, us_match: bool):
    stack_hits = sorted(set(m.group(1).lower() for m in STACK_RE.finditer((text or '').lower())))
    score_cfg = CONFIG['scoring']
    score = int(score_cfg['base_qa']) + len(stack_hits) * int(score_cfg['stack_each']) + (int(score_cfg['us_bonus']) if us_match else 0)
    details = {
        'base_qa': int(score_cfg['base_qa']),
        'stack_hits': stack_hits,
        'stack_points': len(stack_hits) * int(score_cfg['stack_each']),
        'us_bonus': int(score_cfg['us_bonus']) if us_match else 0,
    }
    return score, details



def make_record(title, company, location, url, date_posted='Not specified', min_salary='Not specified', max_salary='Not specified', employment='Not specified', notes=''):
    return {
        'Date Found': time.strftime('%Y-%m-%d'),
        'Job Title': (title or '').strip() or 'Not specified',
        'Company': (company or '').strip() or 'Not specified',
        'Location': normalize_location(location),
        'Job URL': norm(url),
        'Date Posted': (date_posted or '').strip()[:10] if str(date_posted).strip() else 'Not specified',
        'Min Salary': min_salary or 'Not specified',
        'Max Salary': max_salary or 'Not specified',
        'Employment Type': employment or 'Not specified',
        'Notes': notes or '',
    }



def ensure_workbook():
    os.makedirs(RESULTS_DIR, exist_ok=True)
    if os.path.exists(FILE):
        wb = openpyxl.load_workbook(FILE)
        ws = wb.active
        if ensure_schema(ws):
            wb.save(FILE)
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Jobs'
    ws.append(COLUMNS)
    wb.save(FILE)



def load_existing():
    ensure_workbook()
    wb = openpyxl.load_workbook(FILE)
    ws = wb.active
    schema_changed = ensure_schema(ws)
    header = [c.value for c in ws[1]]
    h = {str(v).strip(): i for i, v in enumerate(header) if v}
    rows_changed = hydrate_existing_rows(ws, h)
    if schema_changed or rows_changed:
        wb.save(FILE)

    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(r)]
    existing_url = set()
    existing_tc = set()
    existing_keys = set()
    for r in rows:
        title = str(r[h['Job Title']] or '').strip().lower()
        company = str(r[h['Company']] or '').strip().lower()
        url = norm(str(r[h['Job URL']] or ''))
        location = str(r[h['Location']] or '').strip()
        key = str(r[h['Job Key']] or '').strip() if 'Job Key' in h else ''
        if url:
            existing_url.add(url)
        if key:
            existing_keys.add(key)
        else:
            existing_keys.add(job_key_from_values(url, title, company, location))
        if title and company:
            existing_tc.add((title, company))

    return wb, ws, h, rows, existing_url, existing_tc, existing_keys



def export_json_snapshot():
    wb = openpyxl.load_workbook(FILE)
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not any(r):
            continue
        rows.append({header[i]: (r[i] if i < len(r) else '') for i in range(len(header))})
    with open(JSON_FILE, 'w', encoding='utf-8') as f:
        json.dump(rows, f, ensure_ascii=False, indent=2)
    return len(rows)



def parse_builtin_job(session, url: str):
    ok, _code, final_url, html, _reason = validate_url(session, url)
    if not ok:
        return None
    soup = BeautifulSoup(html, 'lxml')
    ld = soup.find('script', {'type': 'application/ld+json'})
    if not ld:
        return None
    try:
        obj = json.loads(ld.get_text())
    except Exception:
        return None

    jp = None
    for g in obj.get('@graph', []) if isinstance(obj, dict) else []:
        if isinstance(g, dict) and g.get('@type') == 'JobPosting':
            jp = g
            break
    if not jp:
        return None

    title = str(jp.get('title', '')).strip()
    company = str(jp.get('hiringOrganization', {}).get('name', '')).strip() or 'Not specified'
    desc = str(jp.get('description', ''))

    text = f'{title} {desc}'
    if not ROLE_RE.search(text):
        return None

    remote = str(jp.get('jobLocationType', '')).upper() == 'TELECOMMUTE' or ('remote' in text.lower())
    if not remote:
        return None

    us = False
    jl = jp.get('jobLocation')
    locs = jl if isinstance(jl, list) else ([jl] if isinstance(jl, dict) else [])
    for loc in locs:
        addr = loc.get('address', {}) if isinstance(loc, dict) else {}
        c = str(addr.get('addressCountry', '')).upper().strip()
        if c in {'US', 'USA', 'UNITED STATES', 'UNITED STATES OF AMERICA'}:
            us = True
            break

    if not us and not US_RE.search(text.lower()):
        return None

    date_posted = str(jp.get('datePosted', '')).strip()[:10] or 'Not specified'

    min_sal = max_sal = 'Not specified'
    bs = jp.get('baseSalary')
    if isinstance(bs, dict) and isinstance(bs.get('value'), dict):
        cur = bs.get('currency', 'USD')
        mn = bs['value'].get('minValue')
        mx = bs['value'].get('maxValue')
        if mn is not None and mx is not None:
            min_sal, max_sal = f'{cur} {mn}', f'{cur} {mx}'
        elif mn is not None:
            min_sal = max_sal = f'{cur} {mn}'

    et = str(jp.get('employmentType', '')).lower()
    emp = 'Full-time' if 'full' in et else ('Contract' if 'contract' in et else 'Not specified')

    score, details = score_role(text, us)
    notes = 'Source: BuiltIn; Validated URL'
    if details['stack_hits']:
        notes += '; Stack match: ' + ', '.join(details['stack_hits'])
    notes += f"; Score: {score}"

    rec = make_record(
        title=title,
        company=company,
        location='Remote, United States' if us else 'Remote',
        url=final_url or url,
        date_posted=date_posted,
        min_salary=min_sal,
        max_salary=max_sal,
        employment=emp,
        notes=notes,
    )
    rec['_score'] = score
    return rec



def collect_builtin(session):
    terms = BUILTIN_TERMS
    listing_templates = [
        'https://builtin.com/jobs/remote/dev-engineering/qa?page={page}',
    ] + [f'https://builtin.com/jobs/remote/dev-engineering/search/{t}?page={{page}}' for t in terms]

    job_urls = set()
    for tmpl in listing_templates:
        for page in range(1, int(CONFIG['sources']['builtin_max_pages'])):
            u = tmpl.format(page=page)
            r = request_with_retries(session, u)
            if r is None or r.status_code != 200:
                continue
            soup = BeautifulSoup(r.text, 'lxml')
            ld = soup.find('script', {'type': 'application/ld+json'})
            if not ld:
                if page > 2:
                    break
                continue
            try:
                obj = json.loads(ld.get_text())
            except Exception:
                continue
            items = []
            for g in obj.get('@graph', []) if isinstance(obj, dict) else []:
                if isinstance(g, dict) and g.get('@type') == 'ItemList':
                    items = g.get('itemListElement', [])
            if not items and page > 2:
                break
            for it in items:
                if isinstance(it, dict) and it.get('url'):
                    job_urls.add(norm(it['url']))

    out = []
    for u in job_urls:
        job = parse_builtin_job(session, u)
        if job:
            out.append(job)
    return out



def collect_remoteok(session):
    out = []
    r = request_with_retries(session, 'https://remoteok.com/api')
    if r is None:
        return out
    try:
        data = r.json()
    except Exception:
        return out
    for it in data:
        if not isinstance(it, dict) or not it.get('position'):
            continue
        title = str(it.get('position', '')).strip()
        desc = str(it.get('description', '')) + ' ' + ' '.join(it.get('tags', []) if isinstance(it.get('tags'), list) else [])
        text = f'{title} {desc}'
        if not ROLE_RE.search(text):
            continue
        if not REMOTE_RE.search(text):
            continue
        url = norm(str(it.get('url', '')).strip())
        if not url:
            continue
        ok, _code, final_url, _html, _reason = validate_url(session, url)
        if not ok:
            continue
        mn, mx = it.get('salary_min'), it.get('salary_max')
        min_sal, max_sal = (f'USD {mn}', f'USD {mx}') if mn and mx else ('Not specified', 'Not specified')
        score, details = score_role(text, bool(US_RE.search(text.lower())))
        notes = 'Source: RemoteOK; Validated URL'
        if details['stack_hits']:
            notes += '; Stack match: ' + ', '.join(details['stack_hits'])
        notes += f"; Score: {score}"
        rec = make_record(
            title=title,
            company=str(it.get('company', '')).strip() or 'Not specified',
            location='Remote',
            url=final_url or url,
            date_posted='Not specified',
            min_salary=min_sal,
            max_salary=max_sal,
            employment='Not specified',
            notes=notes,
        )
        rec['_score'] = score
        out.append(rec)
    return out



def collect_jobicy(session):
    out = []
    r = request_with_retries(session, 'https://jobicy.com/api/v2/remote-jobs?count=100&geo=usa')
    if r is None:
        return out
    try:
        data = r.json()
    except Exception:
        return out
    for it in data.get('jobs', []):
        title = str(it.get('jobTitle', '')).strip()
        desc = str(it.get('jobDescription', ''))
        text = f'{title} {desc}'
        if not ROLE_RE.search(text):
            continue
        url = norm(str(it.get('url', '')).strip())
        if not url:
            continue
        ok, _code, final_url, _html, _reason = validate_url(session, url)
        if not ok:
            continue
        score, details = score_role(text, bool(US_RE.search(text.lower())))
        notes = 'Source: Jobicy; Validated URL'
        if details['stack_hits']:
            notes += '; Stack match: ' + ', '.join(details['stack_hits'])
        notes += f"; Score: {score}"
        rec = make_record(
            title=title,
            company=str(it.get('companyName', '')).strip() or 'Not specified',
            location=f"Remote ({str(it.get('jobGeo', '')).strip() or 'Worldwide'})",
            url=final_url or url,
            date_posted=str(it.get('pubDate', '')).strip()[:10] or 'Not specified',
            employment='Full-time' if 'full' in str(it.get('jobType', '')).lower() else 'Not specified',
            notes=notes,
        )
        rec['_score'] = score
        out.append(rec)
    return out



def collect_themuse(session):
    out = []
    for page in range(1, int(CONFIG['sources']['themuse_max_pages']) + 1):
        r = request_with_retries(session, f'https://www.themuse.com/api/public/jobs?page={page}')
        if r is None:
            continue
        try:
            data = r.json()
        except Exception:
            continue
        for it in data.get('results', []):
            title = str(it.get('name', '')).strip()
            desc = str(it.get('contents', ''))
            locs = ', '.join([x.get('name', '') for x in it.get('locations', []) if isinstance(x, dict)])
            text = f'{title} {desc} {locs}'
            if not ROLE_RE.search(text):
                continue
            if not looks_like_us_remote(text, locs):
                continue
            refs = it.get('refs', {}) if isinstance(it.get('refs'), dict) else {}
            url = norm(str(refs.get('landing_page', '')).strip())
            if not url:
                continue
            ok, _code, final_url, _html, _reason = validate_url(session, url)
            if not ok:
                continue
            comp = str((it.get('company') or {}).get('name', '')).strip() if isinstance(it.get('company'), dict) else ''
            score, details = score_role(text, True)
            notes = 'Source: TheMuse; Validated URL'
            if details['stack_hits']:
                notes += '; Stack match: ' + ', '.join(details['stack_hits'])
            notes += f"; Score: {score}"
            rec = make_record(
                title=title,
                company=comp or 'Not specified',
                location=locs or 'Remote, United States',
                url=final_url or url,
                date_posted=str(it.get('publication_date', '')).strip()[:10] or 'Not specified',
                notes=notes,
            )
            rec['_score'] = score
            out.append(rec)
    return out



def collect_remotive(session):
    out = []
    r = request_with_retries(session, 'https://remotive.com/api/remote-jobs')
    if r is None:
        return out
    try:
        data = r.json()
    except Exception:
        return out
    for it in data.get('jobs', []):
        title = str(it.get('title', '')).strip()
        desc = str(it.get('description', ''))
        req_loc = str(it.get('candidate_required_location', '')).strip()
        text = f'{title} {desc} {req_loc}'
        if not ROLE_RE.search(text):
            continue
        if not REMOTE_RE.search(text):
            continue
        if not US_RE.search(text.lower()):
            continue
        url = norm(str(it.get('url', '')).strip())
        if not url:
            continue
        ok, _code, final_url, _html, _reason = validate_url(session, url)
        if not ok:
            continue
        min_sal, max_sal = parse_salary_range(str(it.get('salary', '')).strip())
        score, details = score_role(text, True)
        notes = 'Source: Remotive; Validated URL'
        if details['stack_hits']:
            notes += '; Stack match: ' + ', '.join(details['stack_hits'])
        notes += f"; Score: {score}"
        rec = make_record(
            title=title,
            company=str(it.get('company_name', '')).strip() or 'Not specified',
            location=f"Remote ({req_loc})" if req_loc else 'Remote, United States',
            url=final_url or url,
            date_posted=str(it.get('publication_date', '')).strip()[:10] or 'Not specified',
            min_salary=min_sal,
            max_salary=max_sal,
            employment='Contract' if 'contract' in str(it.get('job_type', '')).lower() else 'Not specified',
            notes=notes,
        )
        rec['_score'] = score
        out.append(rec)
    return out



def collect_arbeitnow(session):
    out = []
    for page in range(1, int(CONFIG['sources']['arbeitnow_max_pages']) + 1):
        r = request_with_retries(session, f'https://www.arbeitnow.com/api/job-board-api?page={page}')
        if r is None:
            continue
        try:
            data = r.json()
        except Exception:
            continue
        for it in data.get('data', []):
            title = str(it.get('title', '')).strip()
            desc = str(it.get('description', ''))
            loc = str(it.get('location', '')).strip()
            text = f'{title} {desc} {loc}'
            if not bool(it.get('remote', False)):
                continue
            if not ROLE_RE.search(text):
                continue
            if not US_RE.search(text.lower()):
                continue
            url = norm(str(it.get('url', '')).strip())
            if not url:
                continue
            ok, _code, final_url, _html, _reason = validate_url(session, url)
            if not ok:
                continue
            jtypes = ', '.join(it.get('job_types', []) if isinstance(it.get('job_types'), list) else [])
            score, details = score_role(text, True)
            notes = 'Source: Arbeitnow; Validated URL'
            if details['stack_hits']:
                notes += '; Stack match: ' + ', '.join(details['stack_hits'])
            notes += f"; Score: {score}"
            rec = make_record(
                title=title,
                company=str(it.get('company_name', '')).strip() or 'Not specified',
                location=f"Remote ({loc})" if loc else 'Remote, United States',
                url=final_url or url,
                date_posted=str(it.get('created_at', '')).strip()[:10] or 'Not specified',
                employment='Full-time' if 'full' in jtypes.lower() else 'Not specified',
                notes=notes,
            )
            rec['_score'] = score
            out.append(rec)
    return out



def collect_greenhouse(session):
    out = []
    for board in GREENHOUSE_BOARDS:
        url = f'https://boards-api.greenhouse.io/v1/boards/{board}/jobs'
        r = request_with_retries(session, url)
        if r is None:
            continue
        try:
            data = r.json()
        except Exception:
            continue
        for it in data.get('jobs', []):
            title = str(it.get('title', '')).strip()
            content = str(it.get('content', ''))
            loc = str((it.get('location') or {}).get('name', '')).strip() if isinstance(it.get('location'), dict) else ''
            text = f'{title} {content} {loc}'
            if not ROLE_RE.search(text):
                continue
            if 'remote' not in text.lower():
                continue
            if not US_RE.search(text.lower()):
                continue
            job_url = str(it.get('absolute_url', '')).strip()
            if not job_url:
                continue
            ok, _code, final_url, _html, _reason = validate_url(session, job_url)
            if not ok:
                continue
            score, details = score_role(text, True)
            notes = 'Source: Greenhouse; Validated URL'
            if details['stack_hits']:
                notes += '; Stack match: ' + ', '.join(details['stack_hits'])
            notes += f"; Score: {score}"
            rec = make_record(
                title=title,
                company=board,
                location=loc or 'Remote, United States',
                url=final_url or job_url,
                notes=notes,
            )
            rec['_score'] = score
            out.append(rec)
    return out



def collect_lever(session):
    out = []
    for board in LEVER_BOARDS:
        url = f'https://api.lever.co/v0/postings/{board}?mode=json'
        r = request_with_retries(session, url)
        if r is None:
            continue
        try:
            data = r.json()
        except Exception:
            continue
        if not isinstance(data, list):
            continue
        for it in data:
            title = str(it.get('text', '')).strip()
            desc = str(it.get('descriptionPlain', '') or it.get('description', ''))
            categories = it.get('categories', {}) if isinstance(it.get('categories'), dict) else {}
            loc = str(categories.get('location', '')).strip()
            text = f'{title} {desc} {loc}'
            if not ROLE_RE.search(text):
                continue
            if 'remote' not in text.lower():
                continue
            if not US_RE.search(text.lower()):
                continue
            job_url = str(it.get('hostedUrl', '')).strip()
            if not job_url:
                continue
            ok, _code, final_url, _html, _reason = validate_url(session, job_url)
            if not ok:
                continue
            score, details = score_role(text, True)
            notes = 'Source: Lever; Validated URL'
            if details['stack_hits']:
                notes += '; Stack match: ' + ', '.join(details['stack_hits'])
            notes += f"; Score: {score}"
            rec = make_record(
                title=title,
                company=board,
                location=loc or 'Remote, United States',
                url=final_url or job_url,
                notes=notes,
            )
            rec['_score'] = score
            out.append(rec)
    return out



def dedupe_candidates(cands, existing_url, existing_tc, existing_keys=None):
    existing_keys = existing_keys or set()
    unique = []
    run_u, run_tc, run_k = set(), set(), set()
    stats = {
        'input': len(cands),
        'dropped_empty': 0,
        'dropped_url_existing_or_run': 0,
        'dropped_title_company_existing_or_run': 0,
        'dropped_job_key_existing_or_run': 0,
        'kept': 0,
    }

    for c in sorted(cands, key=lambda x: x.get('_score', 0), reverse=True):
        t = str(c.get('Job Title', '')).strip().lower()
        co = str(c.get('Company', '')).strip().lower()
        lo = str(c.get('Location', '')).strip().lower()
        u = norm(c.get('Job URL', ''))
        k = job_key_from_values(u, t, co, lo)
        tc = (t, co)
        if not t or not co or not u:
            stats['dropped_empty'] += 1
            continue
        if u in existing_url or u in run_u:
            stats['dropped_url_existing_or_run'] += 1
            continue
        if k in existing_keys or k in run_k:
            stats['dropped_job_key_existing_or_run'] += 1
            continue
        if tc in existing_tc or tc in run_tc:
            stats['dropped_title_company_existing_or_run'] += 1
            continue
        c['Job Key'] = k
        unique.append(c)
        run_u.add(u)
        run_tc.add(tc)
        run_k.add(k)

    stats['kept'] = len(unique)
    return unique, stats



def write_run_report(report_path, payload):
    with open(report_path, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)



def active_new_count(rows, h):
    count = 0
    i_app = h.get('Application Status')
    i_link = h.get('Link Status')
    for r in rows:
        app = str(r[i_app] or '').strip().lower() if i_app is not None and i_app < len(r) else ''
        link = str(r[i_link] or '').strip().lower() if i_link is not None and i_link < len(r) else ''
        if app == 'new' and link != 'closed':
            count += 1
    return count


def revalidate_new_openings(session, ws, h):
    i_url = h.get('Job URL')
    i_app = h.get('Application Status')
    i_last = h.get('Last Checked')
    i_link = h.get('Link Status')
    if i_url is None or i_app is None or i_last is None or i_link is None:
        return {'checked': 0, 'active': 0, 'closed': 0, 'blocked': 0, 'unknown': 0}

    stats = {'checked': 0, 'active': 0, 'closed': 0, 'blocked': 0, 'unknown': 0}
    for row_idx in range(2, ws.max_row + 1):
        app = str(ws.cell(row=row_idx, column=i_app + 1).value or '').strip().lower()
        if app != 'new':
            continue

        url = str(ws.cell(row=row_idx, column=i_url + 1).value or '').strip()
        ws.cell(row=row_idx, column=i_last + 1, value=today_str())
        if not url:
            ws.cell(row=row_idx, column=i_link + 1, value='unknown')
            stats['unknown'] += 1
            stats['checked'] += 1
            continue

        ok, _code, _final_url, _html, reason = validate_url(session, url)
        if ok:
            ws.cell(row=row_idx, column=i_link + 1, value='active')
            stats['active'] += 1
        else:
            if reason in {'closed_phrase', 'closed_url_pattern'} or str(reason).startswith('http_4'):
                ws.cell(row=row_idx, column=i_link + 1, value='closed')
                ws.cell(row=row_idx, column=i_app + 1, value='closed')
                stats['closed'] += 1
            elif reason == 'blocked_page':
                ws.cell(row=row_idx, column=i_link + 1, value='blocked')
                stats['blocked'] += 1
            else:
                ws.cell(row=row_idx, column=i_link + 1, value='unknown')
                stats['unknown'] += 1
        stats['checked'] += 1
    return stats


def prune_closed_rows(ws, h):
    i_app = h.get('Application Status')
    i_link = h.get('Link Status')
    if i_app is None and i_link is None:
        return 0

    to_delete = []
    for row_idx in range(2, ws.max_row + 1):
        app = str(ws.cell(row=row_idx, column=i_app + 1).value or '').strip().lower() if i_app is not None else ''
        link = str(ws.cell(row=row_idx, column=i_link + 1).value or '').strip().lower() if i_link is not None else ''
        if app == 'closed' or link == 'closed':
            to_delete.append(row_idx)

    for row_idx in reversed(to_delete):
        ws.delete_rows(row_idx, 1)
    return len(to_delete)


def main():
    global BUILTIN_TERMS
    profile = ensure_profile(PROFILE_PATH, RESUME_DIR)
    runtime_filters = _merge_profile_filters(_filters_from_config(CONFIG), profile)
    apply_runtime_filters(runtime_filters)
    BUILTIN_TERMS = _build_builtin_terms(profile)
    RUN_STATS['profile'] = {
        'profile_path': PROFILE_PATH,
        'resume_dir': RESUME_DIR,
        'resume_file': profile.get('resume_file'),
        'seniority': profile.get('seniority'),
        'role_keywords': profile.get('role_keywords', []),
        'stack_keywords': profile.get('stack_keywords', []),
        'queries': profile.get('queries', []),
    }

    run_mode = os.getenv('RUN_MODE', 'topup').strip().lower()
    log_event('run_start', config_path=CONFIG_PATH, target=TARGET, file=FILE, json_file=JSON_FILE, run_mode=run_mode)
    wb, ws, h, existing_rows, existing_url, existing_tc, existing_keys = load_existing()

    session = requests.Session()
    session.headers.update(HEADERS)

    reval_stats = revalidate_new_openings(session, ws, h)
    RUN_STATS['revalidate'] = reval_stats
    print('REVALIDATE_CHECKED', reval_stats['checked'])
    print('REVALIDATE_ACTIVE', reval_stats['active'])
    print('REVALIDATE_CLOSED', reval_stats['closed'])
    if PRUNE_CLOSED:
        pruned = prune_closed_rows(ws, h)
        RUN_STATS['pruned_closed_rows'] = pruned
        print('PRUNED_CLOSED_ROWS', pruned)
    wb.save(FILE)

    # Re-load after revalidation so queue metrics are current.
    wb, ws, h, existing_rows, existing_url, existing_tc, existing_keys = load_existing()
    current = len(existing_rows)
    active_new = active_new_count(existing_rows, h)
    need = max(0, TARGET - active_new)
    print(f'CURRENT_ROWS {current}')
    print(f'ACTIVE_NEW {active_new}')
    print(f'NEED_ADD {need}')

    if run_mode == 'revalidate':
        cnt = export_json_snapshot()
        print('JSON_WRITTEN', JSON_FILE)
        print('JSON_ROWS', cnt)
        print('REVALIDATE_ONLY_DONE')
        report = {
            **RUN_STATS,
            'run_id': RUN_ID,
            'config_path': CONFIG_PATH,
            'log_path': LOG_PATH,
            'xlsx_path': FILE,
            'json_path': JSON_FILE,
            'final_rows': current,
            'final_active_new': active_new,
            'added_rows': 0,
            'status': 'revalidate_only',
            'run_finished_utc': datetime.now(timezone.utc).isoformat(),
        }
        write_run_report(os.path.join(RESULTS_DIR, 'latest_run_report.json'), report)
        write_run_report(os.path.join(RESULTS_DIR, f'run_report_{RUN_ID}.json'), report)
        return

    if need == 0:
        cnt = export_json_snapshot()
        print('JSON_WRITTEN', JSON_FILE)
        print('JSON_ROWS', cnt)
        print('ALREADY_AT_TARGET')

        report = {
            **RUN_STATS,
            'run_id': RUN_ID,
            'config_path': CONFIG_PATH,
            'log_path': LOG_PATH,
            'xlsx_path': FILE,
            'json_path': JSON_FILE,
            'final_rows': current,
            'final_active_new': active_new,
            'added_rows': 0,
            'status': 'already_at_target',
            'run_finished_utc': datetime.now(timezone.utc).isoformat(),
        }
        write_run_report(os.path.join(RESULTS_DIR, 'latest_run_report.json'), report)
        write_run_report(os.path.join(RESULTS_DIR, f'run_report_{RUN_ID}.json'), report)
        return

    collectors = [
        ('builtin', collect_builtin),
        ('remoteok', collect_remoteok),
        ('jobicy', collect_jobicy),
        ('themuse', collect_themuse),
        ('remotive', collect_remotive),
        ('arbeitnow', collect_arbeitnow),
        ('greenhouse', collect_greenhouse),
        ('lever', collect_lever),
    ]

    cands = []
    for source_name, fn in collectors:
        started = time.time()
        source_rows = fn(session)
        elapsed = round(time.time() - started, 2)
        RUN_STATS['source_counts'][source_name] = len(source_rows)
        cands.extend(source_rows)
        print(f'CAND_{source_name.upper()}', len(source_rows))
        log_event('source_done', source=source_name, rows=len(source_rows), elapsed_seconds=elapsed)

    unique, dedupe_stats = dedupe_candidates(cands, existing_url, existing_tc, existing_keys)
    RUN_STATS['dedupe'] = dedupe_stats

    print('CAND_UNIQUE_NEW', len(unique))
    add = unique[:need]
    print('WILL_ADD', len(add))

    for c in add:
        job_key = c.get('Job Key') or job_key_from_values(
            c.get('Job URL', ''),
            c.get('Job Title', ''),
            c.get('Company', ''),
            c.get('Location', ''),
        )
        ws.append([
            c['Date Found'], c['Job Title'], c['Company'], c['Location'], c['Job URL'],
            c['Date Posted'], c['Min Salary'], c['Max Salary'], c['Employment Type'], c['Notes'],
            'new', '', today_str(), 'active', job_key
        ])

    wb.save(FILE)

    wb2 = openpyxl.load_workbook(FILE)
    ws2 = wb2.active
    header2 = [c.value for c in ws2[1]]
    h2 = {str(v).strip(): i for i, v in enumerate(header2) if v}
    rows2 = [r for r in ws2.iter_rows(min_row=2, values_only=True) if any(r)]
    final_rows = sum(1 for r in ws2.iter_rows(min_row=2, values_only=True) if any(r))
    final_active_new = active_new_count(rows2, h2)
    print('FINAL_ROWS', final_rows)
    print('FINAL_ACTIVE_NEW', final_active_new)

    json_rows = export_json_snapshot()
    print('JSON_WRITTEN', JSON_FILE)
    print('JSON_ROWS', json_rows)

    if final_active_new < TARGET:
        print('SHORTFALL', TARGET - final_active_new)

    report = {
        **RUN_STATS,
        'run_id': RUN_ID,
        'config_path': CONFIG_PATH,
        'log_path': LOG_PATH,
        'xlsx_path': FILE,
        'json_path': JSON_FILE,
        'target': TARGET,
        'starting_rows': current,
        'starting_active_new': active_new,
        'final_rows': final_rows,
        'final_active_new': final_active_new,
        'added_rows': len(add),
        'shortfall': max(0, TARGET - final_active_new),
        'run_finished_utc': datetime.now(timezone.utc).isoformat(),
    }
    write_run_report(os.path.join(RESULTS_DIR, 'latest_run_report.json'), report)
    write_run_report(os.path.join(RESULTS_DIR, f'run_report_{RUN_ID}.json'), report)
    log_event('run_done', final_rows=final_rows, added_rows=len(add), shortfall=max(0, TARGET - final_rows))


if __name__ == '__main__':
    main()
